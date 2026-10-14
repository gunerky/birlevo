"""Independent semantic checks with bundled python-docx/openpyxl. No Word rendering claim."""
import argparse
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from docx import Document
from openpyxl import load_workbook


def verify(native, output, fixtures):
    output.mkdir(parents=True, exist_ok=True)
    doc = Document(native / 'native.docx')
    expected = ['Birlevo basic document', 'İstanbul café العربية فارسی & <text>', 'First\tsecond\nthird']
    assert [p.text for p in doc.paragraphs] == expected
    assert [[cell.text for cell in row.cells] for row in doc.tables[0].rows] == [['Code', 'Text'], ['00123', 'سلام‌دنیا']]
    before = hashlib.sha256((native / 'native.docx').read_bytes()).hexdigest()
    # Deterministic OOXML edit, retaining every other package part; JS document editing surface is not used.
    with ZipFile(native / 'native.docx') as source, ZipFile(output / 'edited.docx', 'w', ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == 'word/document.xml':
                root = ET.fromstring(content)
                root.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t').text = 'Edited document'
                content = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            target.writestr(entry.filename, content)
    reopened = Document(output / 'edited.docx')
    assert reopened.paragraphs[0].text == 'Edited document'
    assert [p.text for p in reopened.paragraphs[1:]] == expected[1:]
    assert before == hashlib.sha256((native / 'native.docx').read_bytes()).hexdigest()
    wb = load_workbook(native / 'native.xlsx', data_only=False)
    assert wb.sheetnames == ['Data']
    assert wb['Data']['A2'].value == '00123'
    assert wb['Data']['A3'].value == '00000'
    assert wb['Data']['A2'].number_format == '@'
    assert wb['Data']['B2'].value == '=2+2'
    assert wb['Data']['B2'].data_type == 's'
    assert wb['Data']['B3'].data_type == 's'
    assert wb['Data']['B4'].value == '0'
    assert not any(c.data_type == 'f' for row in wb['Data'] for c in row)
    source = load_workbook(fixtures / 'rich.xlsx', data_only=False)
    cache = load_workbook(fixtures / 'rich.xlsx', data_only=True)
    assert source['Data']['C2'].value == '=B2*2' and cache['Data']['C2'].value == 20
    assert str(source['Data']['E2'].value.date()) == '2026-10-14'
    missing = load_workbook(fixtures / 'shared-no-cache.xlsx', data_only=True)
    assert missing['Data']['C1'].value is None
    (output / 'verification.json').write_text(json.dumps({
        'independent_docx_text_tables': True, 'docx_edit_roundtrip': True,
        'xlsx_string_types_preserved': True, 'formula_cache_checked_not_recalculated': True,
        'missing_cache_not_invented': True, 'word_layout_rendered': False,
        'microsoft_office_tested': False,
        'xlsx_preview_leading_zeros_verified': False,
        'preview_limitation': 'Independent PNG renderer displays numeric-looking text without leading zeros; stored strings are verified separately.',
    }, indent=2)+'\n', encoding='utf-8')
    print('Independent DOCX paragraphs/tables/edit and XLSX types/cache checks passed.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('native', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('fixtures', type=Path)
    args = parser.parse_args()
    verify(args.native, args.output, args.fixtures)
