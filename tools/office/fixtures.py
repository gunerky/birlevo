"""Synthetic Word and malformed OPC fixtures; bundled host tools only, no user files/network."""
import io
import json
import warnings
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import argparse
from docx import Document
from docx.shared import Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image


def create(directory):
    directory.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph('Birlevo Office compatibility sample', style='Title')
    doc.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph('İstanbul café العربية فارسی & <text>')
    p = doc.add_paragraph()
    p.add_run('Bold').bold = True
    p.add_run(' and italic').italic = True
    doc.add_paragraph('First\tsecond\nthird')
    table = doc.add_table(rows=2, cols=2)
    for r, row in enumerate([['Code', 'Text'], ['00123', 'سلام‌دنیا']]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    doc.sections[0].header.paragraphs[0].text = 'Synthetic header'
    p = doc.add_paragraph()
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to('https://example.invalid/', RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = 'External link text'
    run.append(t); link.append(run); p._p.append(link)
    png = io.BytesIO()
    Image.new('RGB', (32, 32), '#ff625b').save(png, format='PNG')
    png.seek(0)
    doc.add_picture(png, width=Inches(.3))
    doc.save(directory / 'rich.docx')
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        with ZipFile(directory / 'duplicate.zip', 'w', ZIP_DEFLATED) as z:
            z.writestr('same.xml', '<r>one</r>'); z.writestr('same.xml', '<r>two</r>')
    # Separate valid shared-string example with a formula deliberately lacking a cached result.
    ss = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    rr = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    rels = 'http://schemas.openxmlformats.org/package/2006/relationships'
    ct = 'http://schemas.openxmlformats.org/package/2006/content-types'
    parts = {
        '[Content_Types].xml': f'<Types xmlns="{ct}"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
        '_rels/.rels': f'<Relationships xmlns="{rels}"><Relationship Id="rId1" Type="{rr}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        'xl/workbook.xml': f'<workbook xmlns="{ss}" xmlns:r="{rr}"><sheets><sheet name="Data" sheetId="1" r:id="rId1"/></sheets></workbook>',
        'xl/_rels/workbook.xml.rels': f'<Relationships xmlns="{rels}"><Relationship Id="rId1" Type="{rr}/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="{rr}/sharedStrings" Target="sharedStrings.xml"/></Relationships>',
        'xl/sharedStrings.xml': f'<sst xmlns="{ss}" count="1" uniqueCount="1"><si><t>00123</t></si></sst>',
        'xl/worksheets/sheet1.xml': f'<worksheet xmlns="{ss}"><sheetData><row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1"><v>10</v></c><c r="C1"><f>SUM(B1:B2)</f></c></row><row r="2"><c r="B2"><v>5</v></c></row></sheetData></worksheet>',
    }
    parts['[Content_Types].xml'] = parts['[Content_Types].xml'].replace('</Types>', '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/></Types>')
    parts['xl/_rels/workbook.xml.rels'] = parts['xl/_rels/workbook.xml.rels'].replace('</Relationships>', f'<Relationship Id="rId3" Type="{rr}/styles" Target="styles.xml"/></Relationships>')
    parts['xl/styles.xml'] = f'<styleSheet xmlns="{ss}"><cellXfs count="2"><xf numFmtId="0"/><xf numFmtId="14"/></cellXfs></styleSheet>'
    parts['xl/worksheets/sheet1.xml'] = parts['xl/worksheets/sheet1.xml'].replace('</row><row r="2">', '<c r="D1" s="1"><v>46309</v></c></row><row r="2">')
    with ZipFile(directory / 'shared-no-cache.xlsx', 'w', ZIP_DEFLATED) as z:
        for name, value in parts.items(): z.writestr(name, value)
    manifest(directory)


def manifest(directory):
    import hashlib
    records = [{'file': p.name, 'bytes': p.stat().st_size, 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()}
               for p in sorted(directory.iterdir()) if p.suffix in ('.xlsx', '.docx', '.zip')]
    (directory / 'manifest.json').write_text(json.dumps({'synthetic': True, 'fonts_embedded': False, 'files': records}, indent=2)+'\n', encoding='utf-8')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('mode', choices=['create', 'manifest'])
    parser.add_argument('directory', type=Path)
    args = parser.parse_args()
    (create if args.mode == 'create' else manifest)(args.directory)
